"""The robust reader. Every one of these is a file a real export has produced."""

from __future__ import annotations

import pytest

from sot.ingest.reader import (
    DROP_BLANK,
    DROP_NOTES,
    DROP_REPEATED_HEADER,
    DROP_TOTAL,
    classify_non_data,
    clean_cell,
    decode,
    is_notes_row,
    is_null_token,
    read_delimited,
    read_xlsx,
    repair_row,
    sniff_delimiter,
)

HEADER = ["payroll_id", "employee_name", "job_code", "hours_paid"]


# --- encoding ---------------------------------------------------------------


def test_plain_utf8_is_not_a_fallback():
    result = decode("id,name\n1,Peña\n".encode())
    assert result.encoding == "utf-8"
    assert result.used_fallback is False
    assert "Peña" in result.text


def test_bom_is_stripped_without_counting_as_a_fallback():
    result = decode("﻿id,name\n1,x\n".encode("utf-8-sig"))
    assert result.text.startswith("id,name")
    assert result.used_fallback is False


def test_windows_1252_is_decoded_with_accents_intact():
    """M07: a Latin-1 export with an accented name. The name must survive."""
    result = decode("id,name\n1,Peña García-López\n".encode("cp1252"))
    assert result.used_fallback is True
    assert "Peña García-López" in result.text


def test_detection_does_not_get_to_claim_utf16_without_a_bom():
    """Character-set detection will report UTF-16 for a few bytes of Windows-1252 and
    turn a surname into Chinese characters. Wrong-but-plausible is the worst outcome
    for a name, because nobody notices."""
    result = decode("Peña".encode("cp1252"))
    assert result.text == "Peña"
    assert result.encoding == "cp1252"


def test_real_utf16_with_a_bom_is_still_honoured():
    result = decode("id,name\n1,Peña\n".encode("utf-16"))
    assert "Peña" in result.text


def test_decode_never_raises_on_arbitrary_bytes():
    result = decode(bytes(range(256)))
    assert isinstance(result.text, str)


# --- delimiters -------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("a,b,c\n1,2,3\n", ","),
        ("a\tb\tc\n1\t2\t3\n", "\t"),
        ("a;b;c\n1;2;3\n", ";"),
        ("a|b|c\n1|2|3\n", "|"),
    ],
)
def test_delimiter_is_sniffed_from_content(text: str, expected: str):
    assert sniff_delimiter(text)[0] == expected


def test_a_quoted_delimiter_does_not_confuse_the_sniffer():
    """``"LAST, FIRST"`` contains the delimiter and must not make it look inconsistent."""
    text = 'id,name,hours\n1,"REYES, SOFIA",36\n2,"BELL, MARCUS",40\n'
    assert sniff_delimiter(text)[0] == ","


def test_a_delimiter_that_splits_nothing_is_rejected():
    """Semicolons do not appear, so the comma has to win even though both are tried."""
    assert sniff_delimiter("a,b\n1,2\n")[0] == ","


def test_empty_input_gives_a_default_rather_than_raising():
    assert sniff_delimiter("")[0] == ","


def test_crlf_line_endings_read_the_same_as_lf():
    """M06: a BOM and Windows line endings must make no difference at all."""
    lf = read_delimited(b"id,name\n1,x\n2,y\n")
    crlf = read_delimited("﻿id,name\r\n1,x\r\n2,y\r\n".encode("utf-8-sig"))
    assert [row for row in lf.rows if row] == [row for row in crlf.rows if row]


def test_tab_delimited_reads_the_same_as_comma():
    """M08."""
    comma = read_delimited(b"id,name,hours\n1,Ann,36\n")
    tab = read_delimited(b"id\tname\thours\n1\tAnn\t36\n")
    assert comma.rows == tab.rows


# --- null tokens ------------------------------------------------------------


@pytest.mark.parametrize(
    "token", ["", "NA", "n/a", "#N/A", "NULL", "none", "-", "--", "—", "TBD", "?", "  na  "]
)
def test_placeholders_are_recognised_as_missing(token: str):
    assert is_null_token(token)


@pytest.mark.parametrize("value", ["0", "OFF", "N", "A", "36", "no"])
def test_real_values_are_not_mistaken_for_placeholders(value: str):
    assert not is_null_token(value)


def test_clean_cell_trims_and_blanks_placeholders():
    assert clean_cell("  Sofia  ") == "Sofia"
    assert clean_cell(" N/A ") == ""


# --- non-data rows ----------------------------------------------------------


def _classify(row: list[str]) -> str | None:
    return classify_non_data(row, HEADER, numeric_indexes={3}, identifier_indexes={0, 1})


def test_blank_rows_are_dropped():
    assert _classify(["", "", "", ""]) == DROP_BLANK
    assert _classify(["  ", "", "", ""]) == DROP_BLANK


def test_a_repeated_header_is_dropped():
    """Paginated exports reprint the header part way down."""
    assert _classify(list(HEADER)) == DROP_REPEATED_HEADER


def test_a_labelled_total_row_is_dropped():
    assert _classify(["Total", "", "", "116"]) == DROP_TOTAL
    assert _classify(["GRAND TOTAL", "", "", "116"]) == DROP_TOTAL


def test_an_unlabelled_subtotal_is_dropped():
    """Identifiers empty, numbers full: a spreadsheet subtotal with its label removed."""
    assert _classify(["", "", "RN", "116"]) == DROP_TOTAL


def test_a_notes_row_is_dropped():
    assert _classify(["* figures exclude agency staff", "", "", ""]) == DROP_NOTES


def test_a_real_row_survives_all_of_that():
    assert _classify(["P-1", "REYES, SOFIA", "RN", "36"]) is None


def test_one_populated_identifier_is_a_broken_row_not_a_note():
    """It should be quarantined, where somebody can see it, not quietly dropped."""
    assert not is_notes_row(["E201", "", "", ""])


# --- field-count repair -----------------------------------------------------


def test_an_unquoted_name_is_rejoined():
    """M05: the single most common real-world CSV fault in this data."""
    fields = ["P-3011", "LINDQVIST", " ASTRID", "CNA", "RVD"]
    result = repair_row(fields, expected=4, name_indexes=frozenset({1}))
    assert result.repaired
    assert result.fields == ["P-3011", "LINDQVIST, ASTRID", "CNA", "RVD"]
    assert result.merged_value == "LINDQVIST, ASTRID"


def test_repair_only_happens_at_a_column_that_holds_a_name():
    """Without this, a row of codes rejoins just as happily - CNA and RVD look exactly
    like a surname and a given name - and the record is silently corrupted."""
    fields = ["P-1", "36", "CNA", "RVD", "extra"]
    assert not repair_row(fields, expected=4, name_indexes=frozenset({1})).repaired


def test_nothing_is_repaired_when_the_name_column_is_unknown():
    assert not repair_row(["a", "b", "c", "d", "e"], expected=4, name_indexes=frozenset()).repaired


def test_two_extra_fields_are_not_repaired():
    """One stray delimiter is a plausible guess. Two is not, so quarantine it."""
    fields = ["P-1", "A", "B", "C", "RN", "BYS"]
    assert not repair_row(fields, expected=4, name_indexes=frozenset({1})).repaired


def test_a_short_row_is_not_repaired():
    assert not repair_row(["P-1", "A"], expected=4, name_indexes=frozenset({1})).repaired


# --- spreadsheets -----------------------------------------------------------


def test_xlsx_reads_as_a_grid_of_strings():
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["employee_id", "last_name", "hire_date"])
    sheet.append(["E1", "Doe", "2020-03-02"])
    buffer = BytesIO()
    workbook.save(buffer)

    grid = read_xlsx(buffer.getvalue())
    assert grid.rows[0] == ["employee_id", "last_name", "hire_date"]
    assert grid.rows[1] == ["E1", "Doe", "2020-03-02"]
