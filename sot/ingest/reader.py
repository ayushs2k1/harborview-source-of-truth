"""Reading a tabular file without ever refusing to read it.

Principle 1 of the spec: a bad file, page, or row becomes a flag, and everything else
keeps going. That shapes every decision here. The encoding ladder never gives up, the
delimiter is worked out from the content rather than the extension, the header is
looked for rather than assumed to be line one, and a row with the wrong number of
fields is repaired if it plausibly can be and quarantined if it cannot - never
silently dropped, because a quietly missing nurse is worse than a visibly broken row.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from sot.canon import squeeze

#: Delimiters worth trying, in preference order for ties.
DELIMITERS = (",", "\t", ";", "|")

#: Text that means "no value". Compared case-insensitively after trimming.
#: Both dash characters are here because exports disagree about which they use.
NULL_TOKENS = frozenset(
    {
        "",
        "na",
        "n/a",
        "#n/a",
        "n.a.",
        "null",
        "none",
        "nil",
        "-",
        "--",
        "\u2013",
        "\u2014",
        "tbd",
        "?",
    }
)

#: A value that looks like ``LAST, FIRST`` once two stray fields are rejoined.
_NAME_AFTER_MERGE = re.compile(r"^[^,\d]{1,60},\s*[^,\d]{1,60}$")

#: First cell of a totals row.
_TOTAL_CELL = re.compile(r"^(grand\s+|sub\s*)?totals?\b", re.IGNORECASE)

#: First cell of a notes row.
_NOTE_CELL = re.compile(r"^\s*(\*+|notes?\b|#|source:|prepared by)", re.IGNORECASE)

#: Reasons a row is not data. Reported as one aggregated count per file.
DROP_BLANK = "blank"
DROP_REPEATED_HEADER = "repeated header"
DROP_TOTAL = "total"
DROP_NOTES = "notes"


@dataclass(frozen=True)
class DecodeResult:
    text: str
    encoding: str
    used_fallback: bool


@dataclass
class Grid:
    """A file read into rows of raw strings, with how that was achieved."""

    rows: list[list[str]] = field(default_factory=list)
    encoding: str = "utf-8"
    delimiter: str | None = ","
    used_fallback: bool = False
    media_kind: str = "tabular"
    note: str = ""

    def __len__(self) -> int:
        return len(self.rows)


#: Assumed single-byte encoding when a file is not UTF-8. Overridable per pack: it is
#: a statement about where the data comes from, not about how files work.
DEFAULT_SINGLE_BYTE_ENCODING = "cp1252"


def decode(data: bytes, *, single_byte: str = DEFAULT_SINGLE_BYTE_ENCODING) -> DecodeResult:
    """Turn bytes into text, trying the likely encodings in order.

    UTF-8 with a BOM, then plain UTF-8, then detection, then the configured single-byte
    codepage, which maps almost every byte and so cannot really fail. Only a decode
    past plain UTF-8 counts as a fallback worth telling the user about.

    Detection is overruled for single-byte codepages. The Windows codepages differ only
    in which accented letters the high bytes mean, so a short sample cannot distinguish
    them: given a Spanish surname in Windows-1252, character-set detection will cheerfully
    report Windows-1250 and turn an n-tilde into an n-acute. Wrong-but-plausible is the
    worst possible outcome for a name, because nobody notices. The assumed codepage is
    configurable, the fallback is always reported, and the flag tells the reader to
    check any name that looks wrong.
    """
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return DecodeResult(_strip_bom(data.decode(encoding)), "utf-8", used_fallback=False)
        except UnicodeDecodeError:
            continue

    try:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
        if (
            best is not None
            and best.encoding
            and _plausible_encoding(best.encoding, data)
            and not _is_single_byte(best.encoding)
        ):
            return DecodeResult(str(best), best.encoding, used_fallback=True)
    except Exception:
        # Detection is advisory; the fallback below always produces something.
        pass

    return DecodeResult(
        _strip_bom(data.decode(single_byte, errors="replace")), single_byte, used_fallback=True
    )


def _strip_bom(text: str) -> str:
    """Remove a leading byte-order mark however many times it appears.

    ``utf-8-sig`` removes one, which is not always enough: a file re-saved through a
    tool that adds its own leaves two. A surviving mark is invisible on screen and
    silently renames the first column, so the first heading stops matching anything.
    """
    return text.lstrip("\ufeff")


def _is_single_byte(encoding: str) -> bool:
    """True for the eight-bit codepages that detection cannot tell apart."""
    name = encoding.lower().replace("-", "_")
    return name.startswith(("cp12", "cp8", "latin", "iso8859", "iso_8859", "mac_", "koi8"))


#: Byte-order marks that make a wide encoding believable.
_WIDE_BOMS = (
    b"\xff\xfe\x00\x00",
    b"\x00\x00\xfe\xff",
    b"\xff\xfe",
    b"\xfe\xff",
)


def _plausible_encoding(encoding: str, data: bytes) -> bool:
    """Reject a detected encoding that cannot be right for a delimited export.

    Character-set detection is statistical, and on a short or ambiguous sample it will
    happily report UTF-16 for a few bytes of Windows-1252 - turning a surname into
    Chinese characters rather than failing. A wide encoding without a byte-order mark
    is not something a payroll export produces, so it is refused and the cp1252
    fallback is used instead.
    """
    name = encoding.lower().replace("-", "_")
    if name.startswith(("utf_16", "utf_32")):
        return data.startswith(_WIDE_BOMS)
    return True


def sniff_delimiter(text: str, *, sample_lines: int = 40) -> tuple[str, float]:
    """Work out the delimiter from the content.

    Scored by how consistently a candidate splits rows into the same number of fields.
    A candidate that leaves most rows as a single field is not the delimiter, whatever
    the file is called. Parsing goes through ``csv`` so a quoted ``"LAST, FIRST"`` does
    not make the comma look inconsistent.
    """
    lines = [line for line in text.splitlines()[:sample_lines] if line.strip()]
    if not lines:
        return (",", 0.0)

    best_delimiter, best_score, best_width = ",", -1.0, 0
    for delimiter in DELIMITERS:
        try:
            rows = list(csv.reader(lines, delimiter=delimiter))
        except csv.Error:
            continue
        widths = [len(row) for row in rows if row]
        if not widths:
            continue
        modal, occurrences = Counter(widths).most_common(1)[0]
        if modal < 2:
            continue  # it did not split anything, so it is not the delimiter
        score = occurrences / len(widths)
        if (score, modal) > (best_score, best_width):
            best_delimiter, best_score, best_width = delimiter, score, modal

    return (best_delimiter, max(best_score, 0.0))


def read_delimited(data: bytes) -> Grid:
    """Read delimited text into a grid of raw strings."""
    decoded = decode(data)
    delimiter, _ = sniff_delimiter(decoded.text)
    rows = [
        list(row) for row in csv.reader(io.StringIO(decoded.text, newline=""), delimiter=delimiter)
    ]
    return Grid(
        rows=rows,
        encoding=decoded.encoding,
        delimiter=delimiter,
        used_fallback=decoded.used_fallback,
    )


def _excel_cell(value: Any) -> str:
    """Render a spreadsheet cell as the text a CSV export would have contained."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_xlsx(data: bytes) -> Grid:
    """Read the first worksheet of a workbook.

    Accepted as a bonus format. Only the first sheet: a workbook with several sheets of
    different sources is a problem we would rather surface than guess at.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [[_excel_cell(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    note = ""
    if len(workbook.worksheets) > 1:
        note = f"workbook has {len(workbook.worksheets)} sheets; only the first was read"
    return Grid(rows=rows, encoding="xlsx", delimiter=None, media_kind="tabular", note=note)


# ---------------------------------------------------------------------------
# Row classification
# ---------------------------------------------------------------------------


def is_null_token(value: str) -> bool:
    """True when a cell holds a placeholder rather than a value."""
    return squeeze(value).casefold() in NULL_TOKENS


def clean_cell(value: str) -> str:
    """Trim a cell and collapse placeholders to empty."""
    trimmed = squeeze(value)
    return "" if is_null_token(trimmed) else trimmed


def is_blank_row(row: list[str]) -> bool:
    return all(not squeeze(cell) for cell in row)


def is_repeated_header(row: list[str], header: list[str]) -> bool:
    """True when a row is the header printed again, as paginated exports do."""
    from sot.ingest.contracts import header_key

    left = [header_key(cell) for cell in row if squeeze(cell)]
    right = [header_key(cell) for cell in header if squeeze(cell)]
    return bool(left) and left == right


def is_total_row(row: list[str], numeric_indexes: set[int], identifier_indexes: set[int]) -> bool:
    """True for a totals or subtotal line.

    Two shapes: the first cell says so, or the identifying fields are empty while the
    numeric ones are full - which is what a spreadsheet subtotal looks like once the
    label has been dropped.
    """
    cells = [squeeze(cell) for cell in row]
    first = next((cell for cell in cells if cell), "")
    if _TOTAL_CELL.match(first):
        return True
    if not numeric_indexes or not identifier_indexes:
        return False
    identifiers_empty = all(not cells[i] for i in identifier_indexes if i < len(cells))
    numerics_present = any(cells[i] for i in numeric_indexes if i < len(cells))
    return identifiers_empty and numerics_present


def is_notes_row(row: list[str]) -> bool:
    """True for a free-text line that happens to sit in the table.

    Conservative on purpose: one populated cell, and it either reads like prose or is
    explicitly marked. A single populated identifier column is a broken data row, not a
    note, and should be quarantined rather than dropped.
    """
    populated = [(index, squeeze(cell)) for index, cell in enumerate(row) if squeeze(cell)]
    if len(populated) != 1:
        return False
    index, text = populated[0]
    if _NOTE_CELL.match(text):
        return True
    return index == 0 and " " in text and len(text) > 20


def classify_non_data(
    row: list[str],
    header: list[str],
    numeric_indexes: set[int],
    identifier_indexes: set[int],
) -> str | None:
    """Why this row is not data, or ``None`` if it is."""
    if is_blank_row(row):
        return DROP_BLANK
    if is_repeated_header(row, header):
        return DROP_REPEATED_HEADER
    if is_total_row(row, numeric_indexes, identifier_indexes):
        return DROP_TOTAL
    if is_notes_row(row):
        return DROP_NOTES
    return None


# ---------------------------------------------------------------------------
# Field-count repair
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class RepairResult:
    fields: list[str]
    repaired: bool
    merged_at: int | None = None
    merged_value: str = ""


def repair_row(fields: list[str], expected: int, *, name_indexes: frozenset[int]) -> RepairResult:
    """Rejoin a row split by an unquoted comma inside a name.

    Three conditions, all required. Exactly one extra field. The merge lands on a column
    the contract says holds a person's name. And the rejoined value actually reads like
    ``LAST, FIRST``.

    The column condition is the one doing the real work. Without it a row of short
    codes rejoins just as happily, because a pair of three-letter codes is
    indistinguishable from a surname and a given name, and the result is a corrupted
    record that nobody will ever notice. A quarantined row is recoverable; a wrongly
    repaired one is not, so when in doubt this refuses and lets the row be quarantined.
    """
    if len(fields) != expected + 1 or not name_indexes:
        return RepairResult(fields, repaired=False)

    for index in sorted(name_indexes):
        if index >= len(fields) - 1:
            continue
        merged = f"{fields[index]},{fields[index + 1]}"
        if not _NAME_AFTER_MERGE.match(merged.strip()):
            continue
        repaired = [*fields[:index], merged, *fields[index + 2 :]]
        return RepairResult(repaired, repaired=True, merged_at=index, merged_value=merged.strip())

    return RepairResult(fields, repaired=False)
